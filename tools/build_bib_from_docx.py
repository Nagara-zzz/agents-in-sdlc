#!/usr/bin/env python3
"""
扫描 attachments 目录下的 docx 文件，提取元数据并生成 LaTeX .bib 文件（基础版）
用法: python tools\build_bib_from_docx.py "C:\path\to\attachments" latex_src\refs.bib
"""
import sys
from pathlib import Path
from docx import Document
from datetime import datetime


def make_key(name):
    # simple key from filename
    return ''.join(c if c.isalnum() else '_' for c in name)


def extract_meta(docx_path: Path):
    doc = Document(str(docx_path))
    props = doc.core_properties
    title = props.title or ''
    author = props.author or ''
    # attempt to find first non-empty paragraph as title if missing
    if not title:
        for p in doc.paragraphs[:10]:
            t = p.text.strip()
            if len(t) > 3:
                title = t
                break
    # year: try to find a 4-digit year in first 20 paragraphs
    year = ''
    import re
    for p in doc.paragraphs[:20]:
        m = re.search(r'(19|20)\d{2}', p.text)
        if m:
            year = m.group(0)
            break
    return {'title': title or docx_path.stem, 'author': author or 'Unknown', 'year': year or datetime.now().year}


def to_bib_entry(meta, key, filename):
    # Use misc type
    title = meta['title'].replace('"', '"')
    author = meta['author']
    year = meta['year']
    entry = f"@misc{{{key},\n  title = {{{title}}},\n  author = {{{author}}},\n  year = {{{year}}},\n  note = {{Converted from {filename}}},\n}}\n\n"
    return entry


def main():
    if len(sys.argv) < 3:
        print('Usage: build_bib_from_docx.py ATTACHMENTS_DIR OUT_BIB', file=sys.stderr)
        sys.exit(2)
    attach_dir = Path(sys.argv[1])
    out_bib = Path(sys.argv[2])
    if not attach_dir.exists():
        print('Attachments dir not found:', attach_dir, file=sys.stderr)
        sys.exit(3)
    out_bib.parent.mkdir(parents=True, exist_ok=True)
    entries = []
    for p in sorted(attach_dir.glob('*.docx')):
        meta = extract_meta(p)
        key = make_key(p.stem)
        entries.append(to_bib_entry(meta, key, p.name))
    out_bib.write_text(''.join(entries), encoding='utf-8')
    print('Wrote', out_bib)

if __name__ == '__main__':
    main()
